"""
generate_cover_letter.py — Exporta una cover letter (Markdown → PDF + DOCX). DETERMINÍSTICO.

La redacción (tono, gancho, adaptación a la empresa) la hace Claude Code en la conversación
— NO consume Anthropic API tokens.

Flujo:
1. Yo (Claude Code) redacto la cover letter en Markdown (≤350 palabras) y la guardo en
   output/cover_letters/<position_id>/cover_letter.md.
2. `--export` la convierte a PDF + DOCX.
3. Se sube a Drive vía MCP (lo hago yo como agente).

Uso:
  python core/generate_cover_letter.py --export --position-id <id>
  python core/generate_cover_letter.py --export --position-id <id> --md-file ruta.md
  cat cl.md | python core/generate_cover_letter.py --export --position-id <id> --stdin
"""

import argparse
import re
import sys
from pathlib import Path

OUTPUT_DIR = Path("output/cover_letters")


def safe_id(position_id: str) -> str:
    return re.sub(r"[^a-z0-9\-]", "", position_id.lower().replace(" ", "-"))[:50]


def markdown_to_pdf(md_content: str, output_path: Path) -> bool:
    try:
        import markdown
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("markdown o playwright no instalados. Instalá con: pip install markdown playwright")
        return False
    try:
        html_content = markdown.markdown(md_content, extensions=["extra"])
        styled_html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
  /* Harvard-style header to match the CV, letter-style body */
  body {{ font-family: Georgia, 'Times New Roman', serif; font-size: 10.5pt; line-height: 1.5; margin: 0; color: #000; }}
  h1 {{ font-size: 19pt; text-align: center; margin: 0 0 2px; letter-spacing: 0.5px; }}
  h1 + p {{ text-align: center; margin: 0 0 10px; font-size: 10pt; }}
  p {{ margin: 8px 0; }}
  a {{ color: #000; text-decoration: underline; }}
  hr {{ border: none; border-top: 1px solid #000; margin: 10px 0; }}
</style></head><body>{html_content}</body></html>"""
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(styled_html, wait_until="load")
            page.pdf(
                path=str(output_path),
                format="A4",
                print_background=True,
                margin={"top": "2.2cm", "bottom": "2.2cm", "left": "2.2cm", "right": "2.2cm"},
            )
            browser.close()
        return True
    except Exception as e:
        print(f"Error generando PDF: {e}")
        return False


_MD_INLINE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|\*\*(.+?)\*\*")


def _add_docx_hyperlink(paragraph, url: str, text: str):
    """Agrega un hipervínculo real (clickable) a un párrafo de python-docx."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.shared import OxmlElement, qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")  # links en negro, como el PDF (estilo Harvard)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _add_docx_runs(paragraph, text: str):
    """Escribe `text` interpretando **negrita** y [texto](url) como hipervínculos
    clickables (antes el DOCX los dejaba en crudo)."""
    pos = 0
    for m in _MD_INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(1) is not None:  # [texto](url)
            _add_docx_hyperlink(paragraph, m.group(2), m.group(1))
        else:  # **negrita**
            paragraph.add_run(m.group(3)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _docx_bottom_border(paragraph):
    """Línea inferior de un párrafo (separador estilo Harvard, igual que el PDF)."""
    from docx.oxml.shared import OxmlElement, qn
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    pPr.append(pbdr)


def _docx_set_base_font(doc, name: str = "Georgia", size_pt: float = 10.5):
    """Fija la fuente serif base (Normal) para que TODO el DOCX use Georgia como el PDF."""
    from docx.shared import Pt
    from docx.oxml.ns import qn
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


def markdown_to_docx(md_content: str, output_path: Path) -> bool:
    """DOCX con el MISMO formato Harvard que el PDF (serif Georgia, nombre centrado,
    encabezado centrado con hipervínculos, cuerpo de carta) y editable."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx no instalado. Instalá con: pip install python-docx")
        return False
    try:
        doc = Document()
        _docx_set_base_font(doc, "Georgia", 10.5)
        for section in doc.sections:
            section.top_margin = Inches(0.87)
            section.bottom_margin = Inches(0.87)
            section.left_margin = Inches(0.87)
            section.right_margin = Inches(0.87)

        in_header = True  # nombre + contacto (antes del '---') van centrados

        for raw in md_content.split("\n"):
            line = raw.rstrip()
            if line.startswith("# "):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(line[2:].strip())
                run.bold = True
                run.font.size = Pt(19)
                in_header = True
            elif line.startswith("## "):
                p = doc.add_paragraph()
                run = p.add_run(line[3:].strip().upper())
                run.bold = True
                run.font.size = Pt(11)
                _docx_bottom_border(p)
                in_header = False
            elif line.strip() == "---":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                _docx_bottom_border(p)
                in_header = False
            elif line.strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                _add_docx_runs(p, line.strip())
                if in_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(10)

        doc.save(str(output_path))
        return True
    except Exception as e:
        print(f"Error generando DOCX: {e}")
        return False


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")


def doc_basename(candidate: str, doctype: str, role: str) -> str:
    """Nombre de archivo estilo Nombre-Apellido-TipoDocumento-RolObjetivo."""
    parts = [_slug(candidate), doctype, _slug(role)]
    return "-".join(p for p in parts if p)


def export(position_id: str, md_content: str, candidate: str = "Candidate",
           role: str = "") -> dict:
    cl_dir = OUTPUT_DIR / safe_id(position_id)
    cl_dir.mkdir(parents=True, exist_ok=True)

    base = doc_basename(candidate, "CoverLetter", role) if role else "cover_letter"
    md_path = cl_dir / "cover_letter.md"          # fuente de trabajo (nombre fijo)
    pdf_path = cl_dir / f"{base}.pdf"             # entregable con nombre-convención
    docx_path = cl_dir / f"{base}.docx"

    md_path.write_text(md_content, encoding="utf-8")
    print(f"Cover letter Markdown: {md_path}")

    pdf_ok = markdown_to_pdf(md_content, pdf_path)
    if pdf_ok:
        print(f"Cover letter PDF:      {pdf_path}")
    docx_ok = markdown_to_docx(md_content, docx_path)
    if docx_ok:
        print(f"Cover letter DOCX:     {docx_path}")

    return {
        "md": str(md_path),
        "pdf": str(pdf_path) if pdf_ok else None,
        "docx": str(docx_path) if docx_ok else None,
    }


def main():
    parser = argparse.ArgumentParser(description="Exporta una cover letter Markdown a PDF + DOCX")
    parser.add_argument("--export", action="store_true", help="Exporta el Markdown a PDF + DOCX")
    parser.add_argument("--position-id", required=True, help="ID de la posición (define la carpeta)")
    parser.add_argument("--md-file", help="Ruta al Markdown (default: output/cover_letters/<id>/cover_letter.md)")
    parser.add_argument("--stdin", action="store_true", help="Lee el Markdown desde stdin")
    parser.add_argument("--role", default="", help="Rol objetivo para el nombre del archivo (ej. 'AI Workflow Engineer'). "
                                                    "Genera Nombre-Apellido-CoverLetter-Rol.pdf/.docx")
    parser.add_argument("--candidate", default="Candidate", help="Nombre del candidato para el nombre del archivo")
    args = parser.parse_args()

    if not args.export:
        parser.print_help()
        return

    if args.stdin:
        md_content = sys.stdin.read()
    elif args.md_file:
        md_content = Path(args.md_file).read_text(encoding="utf-8")
    else:
        default_md = OUTPUT_DIR / safe_id(args.position_id) / "cover_letter.md"
        if not default_md.exists():
            print(f"No se encontró {default_md}. Pasá --md-file o --stdin, o escribí primero el .md.")
            sys.exit(1)
        md_content = default_md.read_text(encoding="utf-8")

    if not md_content.strip():
        print("Error: el Markdown de la cover letter está vacío.")
        sys.exit(1)

    export(args.position_id, md_content, candidate=args.candidate, role=args.role)
    print("\nCover letter lista. Subila a Drive con core/upload_to_drive.py.")


if __name__ == "__main__":
    main()
